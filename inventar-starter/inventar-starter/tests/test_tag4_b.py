"""Export-Tests: python -m unittest discover -s tests -p test_tag4_b.py -v.

Zusätzlich zu requirements.txt: httpx==0.28.1 und openpyxl==3.1.5.
"""

import codecs
import csv
from datetime import datetime, timedelta, timezone
from io import BytesIO, StringIO
import sqlite3
import unittest
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app import main
from backend.app.main import app


class ExportTests(unittest.TestCase):
    def setUp(self):
        self.rows = [dict(zip(main.EXPORT_COLUMNS, (
            "000123", "Laptop", "Gebäude E", "Ben König",
            datetime(2026, 10, 25, 10, 30, tzinfo=timezone(timedelta(hours=2))),
            None, True, None,
        ))), dict(zip(main.EXPORT_COLUMNS, (
            "LAP-004", "Laptop", "Gebäude F", "Clara Meier",
            datetime(2026, 10, 24, 9, tzinfo=timezone.utc),
            datetime(2026, 10, 25, 8, tzinfo=timezone.utc),
            False, 'Gehäuse; "Kratzer"\nScharnier prüfen',
        )))]

    def test_csv_headers_unicode_quoting_dates_and_nulls(self):
        with patch.object(main, "assignment_rows", return_value=self.rows), TestClient(app) as client:
            response = client.get("/assignments.csv")
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(codecs.BOM_UTF8))
        self.assertIn("text/csv", response.headers["content-type"])
        self.assertIn('filename="assignments.csv"', response.headers["content-disposition"])
        parsed = list(csv.DictReader(StringIO(response.content.decode("utf-8-sig")), delimiter=";"))
        self.assertEqual(tuple(parsed[0]), main.EXPORT_COLUMNS)
        self.assertTrue({"assignment_id", "device_id", "person_id"}.isdisjoint(parsed[0]))
        self.assertEqual(len(parsed[0]), 8)
        self.assertEqual(parsed[0]["serial_number"], "000123")
        self.assertEqual(parsed[0]["location_name"], "Gebäude E")
        self.assertEqual(parsed[0]["issued_at"], "2026-10-25T08:30:00+00:00")
        self.assertEqual(parsed[0]["returned_at"], "")
        self.assertEqual(parsed[0]["is_active"], "true")
        self.assertEqual(parsed[1]["is_active"], "false")
        self.assertEqual(parsed[1]["damage_note"], self.rows[1]["damage_note"])

    def test_xlsx_preserves_cell_types_and_formulas_as_text(self):
        self.rows[0]["damage_note"] = "=1+1"
        with patch.object(main, "assignment_rows", return_value=self.rows), TestClient(app) as client:
            response = client.get("/assignments.xlsx")
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml.sheet", response.headers["content-type"])
        self.assertIn('filename="assignments.xlsx"', response.headers["content-disposition"])
        book = load_workbook(BytesIO(response.content))
        self.assertEqual(book.sheetnames, ["Assignments"])
        sheet = book["Assignments"]
        self.assertEqual(tuple(c.value for c in sheet[1]), main.EXPORT_COLUMNS)
        self.assertEqual(sheet.max_column, 8)
        self.assertTrue({"assignment_id", "device_id", "person_id"}.isdisjoint(c.value for c in sheet[1]))
        self.assertEqual(sheet["A2"].value, "000123")
        self.assertEqual(sheet["A2"].data_type, "s")
        self.assertEqual(sheet["E2"].value, datetime(2026, 10, 25, 8, 30))
        self.assertEqual(sheet["E2"].data_type, "d")
        self.assertIsNone(sheet["F2"].value)
        self.assertIs(sheet["G2"].value, True)
        self.assertEqual(sheet["H2"].value, "=1+1")
        self.assertEqual(sheet["H2"].data_type, "s")
        self.assertEqual(sheet["H3"].value, self.rows[1]["damage_note"])
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet.auto_filter.ref, "A1:H3")
        book.close()

    def test_empty_exports_still_contain_all_headers(self):
        with patch.object(main, "assignment_rows", return_value=[]), TestClient(app) as client:
            csv_response = client.get("/assignments.csv")
            xlsx_response = client.get("/assignments.xlsx")
        records = list(csv.reader(StringIO(csv_response.content.decode("utf-8-sig")), delimiter=";"))
        self.assertEqual(records, [list(main.EXPORT_COLUMNS)])
        book = load_workbook(BytesIO(xlsx_response.content))
        sheet = book["Assignments"]
        self.assertEqual(sheet.max_row, 1)
        self.assertEqual(tuple(c.value for c in sheet[1]), main.EXPORT_COLUMNS)
        book.close()

    def test_csv_formula_like_input_is_escaped(self):
        self.rows[0]["damage_note"] = "  =1+1"
        with patch.object(main, "assignment_rows", return_value=self.rows), TestClient(app) as client:
            response = client.get("/assignments.csv")
        records = list(csv.DictReader(StringIO(response.content.decode("utf-8-sig")), delimiter=";"))
        self.assertEqual(records[0]["damage_note"], "'  =1+1")

    def test_query_joins_entities_and_keeps_all_assignments_in_order(self):
        with sqlite3.connect(":memory:") as conn:
            conn.row_factory = sqlite3.Row
            conn.executescript("""
                CREATE TABLE device_type (device_type_id INTEGER, name TEXT);
                CREATE TABLE location (location_id INTEGER, name TEXT);
                CREATE TABLE person (person_id INTEGER, first_name TEXT, last_name TEXT);
                CREATE TABLE device (device_id INTEGER, serial_number TEXT, device_type_id INTEGER, location_id INTEGER);
                CREATE TABLE assignment (assignment_id INTEGER, device_id INTEGER, person_id INTEGER, issued_at TEXT, returned_at TEXT, damage_note TEXT);
                INSERT INTO device_type VALUES (8, 'Monitor');
                INSERT INTO location VALUES (9, 'Gebäude H');
                INSERT INTO person VALUES (4, 'Anna', 'Koch');
                INSERT INTO device VALUES (2, 'MON-005', 8, 9);
                INSERT INTO assignment VALUES (5, 2, 4, '2026-10-24', '2026-10-25', 'Kratzer');
                INSERT INTO assignment VALUES (6, 2, 4, '2026-10-25', NULL, 'Zweite Ausgabe');
                INSERT INTO assignment VALUES (7, 2, 4, '2026-10-25', NULL, 'Dritte Ausgabe');
            """)
            result = [dict(row) for row in conn.execute(main.ASSIGNMENT_SQL)]
        self.assertEqual([row["damage_note"] for row in result], ["Dritte Ausgabe", "Zweite Ausgabe", "Kratzer"])
        self.assertEqual(tuple(result[0]), main.EXPORT_COLUMNS)
        self.assertTrue({"assignment_id", "device_id", "person_id"}.isdisjoint(result[0]))
        self.assertEqual(result[0]["person_name"], "Anna Koch")
        self.assertEqual(result[0]["device_type"], "Monitor")
        self.assertEqual(result[0]["location_name"], "Gebäude H")
        self.assertTrue(result[0]["is_active"])
        self.assertFalse(result[2]["is_active"])


if __name__ == "__main__":
    unittest.main()
